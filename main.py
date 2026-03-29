from openpyxl import load_workbook
import requests
import time

TOKEN = "8335948027:AAFDSemc1qjT9rBdLnGsYUkYdGnMMWcpkGY"
URL = f"https://api.telegram.org/bot{TOKEN}/"

user_mode = {}


def main_menu():
    return {
        "keyboard": [
            ["🔧 По машине", "📦 По форме"],
            ["📏 По дорну", "🔪 По ножу"]
        ],
        "resize_keyboard": True,
        "one_time_keyboard": False
    }


def search_xlsx(query, mode):
    wb = load_workbook("load_workbook("data.xlsx")
    sheet = wb.active

    results = []

    for row in sheet.iter_rows(values_only=True):
        if not any(row):
            continue

        machine = str(row[0]) if len(row) > 0 else ""
        volume = str(row[1]) if len(row) > 1 else ""
        name = str(row[2]) if len(row) > 2 else ""
        dorn = str(row[3]) if len(row) > 3 else ""
        soplo = str(row[4]) if len(row) > 4 else ""
        knife = str(row[5]) if len(row) > 5 else ""

        if mode == "machine":
            field = machine
        elif mode == "name":
            field = name
        elif mode == "dorn":
            field = dorn
        elif mode == "knife":
            field = knife
        else:
            field = " ".join([machine, name, dorn, knife])

        if query.lower() in str(field).lower():
            formatted = (
                f"🔧 Машина: {machine}\n"
                f"📦 Литраж: {volume}\n"
                f"📦 Форма: {name}\n\n"
                f"📏 Дорн: {dorn}\n"
                f"🔥 Сопло: {soplo}\n"
                f"🔪 Нож: {knife}"
            )

            results.append(formatted)

    return results if results else ["Ничего не найдено"]


def send_message(chat_id, text):
    requests.post(URL + "sendMessage", json={
        "chat_id": chat_id,
        "text": text,
        "reply_markup": main_menu()
    })


def get_updates(offset=None):
    return requests.get(URL + "getUpdates", params={"offset": offset}).json()


offset = None

while True:
    updates = get_updates(offset)

    for update in updates.get("result", []):
        offset = update["update_id"] + 1

        if "message" not in update:
            continue

        chat_id = update["message"]["chat"]["id"]
        text = update["message"].get("text", "")

        if text == "/start":
            user_mode[chat_id] = None
            send_message(chat_id, "Выбери тип поиска:")

        elif text == "🔧 По машине":
            user_mode[chat_id] = "machine"
            send_message(chat_id, "Введи номер машины:")

        elif text == "📦 По форме":
            user_mode[chat_id] = "name"
            send_message(chat_id, "Введи форму:")

        elif text == "📏 По дорну":
            user_mode[chat_id] = "dorn"
            send_message(chat_id, "Введи значение дорна:")

        elif text == "🔪 По ножу":
            user_mode[chat_id] = "knife"
            send_message(chat_id, "Введи значение ножа:")

        else:
            mode = user_mode.get(chat_id)

            if not mode:
                send_message(chat_id, "Сначала выбери тип поиска")
                continue

            results = search_xlsx(text, mode)
            send_message(chat_id, "\n\n".join(results))

    time.sleep(1)